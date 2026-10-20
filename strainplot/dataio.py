"""数据读取模块
=================

从用户选择的文件夹中自动提取"拉伸应变 / 拉伸应力"两列数据。

支持的数据源
------------
* ``.csv`` / ``.txt`` / ``.dat``：逗号 / 制表符 / 分号分隔的文本表格（自动识别 UTF-8、
  UTF-16、GBK 编码）。适用于 Instron Bluehill 的 ``*_RawData\\Specimen_RawData_*.csv``。
* ``.xlsx`` / ``.xlsm``：Excel 表格。兼容"表头行 + 单位行 + 样品名行"这种多行表头，
  以及普通的单行表头。适用于 H023.xlsx 这类"多组 Stress/Strain 成对排布"的文件。

列匹配规则（不区分大小写，忽略空格与括号）
--------------------------------------------
应变列：包含 ``应变`` / ``strain`` ；应力列：包含 ``应力`` / ``stress``。
若列名同时出现"拉伸"，优先级更高。
"""
from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from typing import Iterable, Sequence

import numpy as np

TEXT_EXT = {".csv", ".txt", ".dat", ".tsv"}
EXCEL_EXT = {".xlsx", ".xlsm"}
SKIP_DIR_HINTS = ("__macosx",)

_STRAIN_KEYS = ("应变", "strain", "elongation")
_STRESS_KEYS = ("应力", "stress")
_UNIT_RE = re.compile(r"^\(?\s*[%a-zA-Zμµ°²³/·\.]+\s*\)?$")


# --------------------------------------------------------------------------- #
# 数据结构
# --------------------------------------------------------------------------- #
@dataclass
class Series:
    """一条曲线（一个样品的应力-应变数据）。"""

    key: str
    label: str
    group: str
    source: str
    x: np.ndarray
    y: np.ndarray
    x_unit: str = "%"
    y_unit: str = "kPa"
    kind: str = "csv"
    selected: bool = True

    # ---- 派生量 -----------------------------------------------------------
    @property
    def n(self) -> int:
        return int(self.x.size)

    @property
    def x_max(self) -> float:
        return float(np.nanmax(self.x)) if self.n else 0.0

    @property
    def y_max(self) -> float:
        return float(np.nanmax(self.y)) if self.n else 0.0

    def summary(self) -> str:
        if not self.n:
            return "空"
        return f"n={self.n}  max ε={self.x_max:g}{self.x_unit}  max σ={self.y_max:g}{self.y_unit}"


@dataclass
class ScanResult:
    series: list[Series] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    def add_warning(self, msg: str) -> None:
        if msg not in self.warnings:
            self.warnings.append(msg)


# --------------------------------------------------------------------------- #
# 基础工具
# --------------------------------------------------------------------------- #
def decode_bytes(raw: bytes) -> str:
    """尽量无损地把表格文件解码成字符串。"""
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig", errors="replace")
    for enc in ("utf-8", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def to_float(cell) -> float | None:
    """把任意单元格内容解析成 float，失败返回 None。"""
    if cell is None:
        return None
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        v = float(cell)
        return v if np.isfinite(v) else None
    s = str(cell).strip().strip('"').strip("'").replace("\u2212", "-")
    if not s:
        return None
    s = s.replace(" ", "")
    try:
        v = float(s)
    except ValueError:
        # 千分位 1,234.5 或 小数逗号 1,5
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s):
            s = s.replace(",", "")
            try:
                v = float(s)
            except ValueError:
                return None
        elif re.fullmatch(r"-?\d+,\d+", s):
            s = s.replace(",", ".")
            try:
                v = float(s)
            except ValueError:
                return None
        else:
            return None
    return v if np.isfinite(v) else None


def _norm(text) -> str:
    s = str(text or "").lower()
    for ch in " \t\r\n_-·.":
        s = s.replace(ch, "")
    return s


def _is_strain(name) -> bool:
    n = _norm(name)
    return any(k in n for k in _STRAIN_KEYS)


def _is_stress(name) -> bool:
    n = _norm(name)
    return any(k in n for k in _STRESS_KEYS)


def _score_strain(name) -> int:
    n = _norm(name)
    score = 0
    if "拉伸" in n:
        score += 4
    if "应变" in n:
        score += 3
    if "strain" in n:
        score += 3
    if "位移" in n:
        score += 1
    if "时间" in n or "time" in n:
        score -= 5
    return score


def _score_stress(name) -> int:
    n = _norm(name)
    score = 0
    if "拉伸" in n:
        score += 4
    if "应力" in n:
        score += 3
    if "stress" in n:
        score += 3
    if "载荷" in n or "load" in n:
        score -= 5
    return score


def _clean_unit(text) -> str:
    s = str(text or "").strip().strip('"')
    s = s.strip("()（）").strip()
    return s


def _guess_unit(unit_cell, fallback: str) -> str:
    u = _clean_unit(unit_cell)
    return u if u and len(u) <= 10 else fallback


def _clean_label(stem: str) -> str:
    s = stem.strip()
    low = s.lower()
    for suffix in (".is_tens_rawdata", ".is_comp_rawdata", "_rawdata", ".is_tens", ".is_comp",
                   ".id_tens"):
        while low.endswith(suffix):
            s = s[: -len(suffix)]
            low = s.lower()
    s = re.sub(r"(?i)^specimen[_\-\s]?rawdata[_\-\s]?", "", s)
    s = s.replace("  ", " ").strip(" -")
    return s or stem


# --------------------------------------------------------------------------- #
# 文本表格（CSV / TXT）
# --------------------------------------------------------------------------- #
def _split_lines(text: str) -> list[list[str]]:
    sample = text[:8000]
    delim = ","
    best = -1
    for cand in (",", "\t", ";", "|"):
        cnt = sample.count(cand)
        if cnt > best:
            best, delim = cnt, cand
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = []
    for row in reader:
        rows.append([c.strip() for c in row])
        if len(rows) > 500000:
            break
    return rows


def _find_header(rows: Sequence[Sequence[str]], max_scan: int = 12) -> int:
    best_idx, best_score = -1, 0
    for i, row in enumerate(rows[:max_scan]):
        has_strain = any(_is_strain(c) for c in row)
        has_stress = any(_is_stress(c) for c in row)
        if has_strain and has_stress:
            return i
        score = int(has_strain) + int(has_stress)
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx if best_score > 0 else -1


def parse_text_table(path: str) -> list[Series]:
    """解析一个文本表格，返回其中的曲线（一般 1 条）。"""
    with open(path, "rb") as fh:
        raw = fh.read()
    text = decode_bytes(raw)
    rows = _split_lines(text)
    if len(rows) < 2:
        raise ValueError("文件内容为空")

    hdr = _find_header(rows)
    if hdr < 0:
        raise ValueError("找不到包含“应变 / 应力”的表头行")

    header = rows[hdr]
    strain_cols = [i for i, c in enumerate(header) if _is_strain(c)]
    stress_cols = [i for i, c in enumerate(header) if _is_stress(c)]
    if not strain_cols or not stress_cols:
        raise ValueError("表头中缺少应变列或应力列")

    xi = max(strain_cols, key=lambda i: _score_strain(header[i]))
    yi = max(stress_cols, key=lambda i: _score_stress(header[i]))
    if xi == yi:
        raise ValueError("应变列与应力列相同")

    # 单位行：(kPa),(%) 这种形式
    xu, yu = "%", "kPa"
    if hdr + 1 < len(rows):
        probe = rows[hdr + 1]
        if probe and not any(to_float(c) is not None for c in probe[: max(xi, yi) + 1]):
            xu = _guess_unit(probe[xi] if xi < len(probe) else "", xu)
            yu = _guess_unit(probe[yi] if yi < len(probe) else "", yu)

    xs, ys = [], []
    for row in rows[hdr + 1:]:
        if len(row) <= max(xi, yi):
            continue
        xv, yv = to_float(row[xi]), to_float(row[yi])
        if xv is None or yv is None:
            continue
        xs.append(xv)
        ys.append(yv)
    if len(xs) < 2:
        raise ValueError("有效数据点少于 2 个")

    parent = os.path.basename(os.path.dirname(os.path.abspath(path)))
    stem = os.path.splitext(os.path.basename(path))[0]
    group = _clean_label(parent)
    label = stem
    if parent.lower().endswith(("_rawdata", "_raw")):
        token = re.sub(r"(?i)^specimen[_\-\s]?rawdata[_\-\s]?", "", stem)
        token = token or stem
        label = f"{group} #{token}"
    label = _clean_label(label)

    return [
        Series(
            key=path,
            label=label,
            group=group,
            source=path,
            x=np.asarray(xs, float),
            y=np.asarray(ys, float),
            x_unit=xu,
            y_unit=yu,
            kind="csv",
        )
    ]


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _row_texts(row) -> list[str]:
    return ["" if c is None else (c if isinstance(c, str) else str(c)) for c in row]


def _is_number_cell(cell) -> bool:
    return to_float(cell) is not None and not isinstance(cell, bool)


def parse_excel(path: str) -> list[Series]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[Series] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for r_i, row in enumerate(ws.iter_rows(values_only=True)):
                rows.append(list(row))
                if r_i > 200000:
                    break
            if len(rows) < 2:
                continue
            out.extend(_series_from_sheet(rows, ws.title, path))
    finally:
        wb.close()
    if not out:
        raise ValueError("Excel 中没有找到 Stress/Strain 数据块")
    return out


def _series_from_sheet(rows: list[list], sheet: str, path: str) -> list[Series]:
    header_idx = -1
    for i, row in enumerate(rows[:20]):
        texts = _row_texts(row)
        if any(_is_strain(c) for c in texts) and any(_is_stress(c) for c in texts):
            header_idx = i
            break
    if header_idx < 0:
        return []

    header = _row_texts(rows[header_idx])
    ncol = len(header)

    # 判断单位行 / 样品名行
    unit_row: list[str] | None = None
    name_row: list[str] | None = None
    data_start = header_idx + 1
    for k in range(header_idx + 1, min(header_idx + 4, len(rows))):
        texts = _row_texts(rows[k])
        non_empty = [t for t in texts if t.strip()]
        numeric = [t for t in non_empty if to_float(t) is not None]
        if not non_empty:
            data_start = k + 1
            continue
        if len(numeric) >= max(1, len(non_empty) * 0.6):
            data_start = k
            break
        if unit_row is None and all(
            (not t.strip()) or _UNIT_RE.match(t.strip()) or len(t.strip()) <= 6 for t in non_empty
        ):
            unit_row = texts
            data_start = k + 1
            continue
        if name_row is None:
            name_row = texts
            data_start = k + 1
            continue
        data_start = k
        break

    pairs: list[tuple[int, int]] = []  # (strain_col, stress_col)
    used = set()
    for i, h in enumerate(header):
        if _is_stress(h) and i not in used:
            cand = [i + 1, i - 1, i + 2, i - 2]
            for c in cand:
                if 0 <= c < ncol and _is_strain(header[c]) and c not in used:
                    pairs.append((c, i))
                    used.add(c)
                    used.add(i)
                    break
    if not pairs:
        return []

    result: list[Series] = []
    for order, (xi, yi) in enumerate(pairs):
        xs, ys = [], []
        for row in rows[data_start:]:
            if len(row) <= max(xi, yi):
                continue
            xv, yv = to_float(row[xi]), to_float(row[yi])
            if xv is None or yv is None:
                continue
            xs.append(xv)
            ys.append(yv)
        if len(xs) < 2:
            continue
        label = ""
        for src in (name_row, header):
            if src and xi < len(src) and str(src[xi] or "").strip():
                label = str(src[xi]).strip()
                break
        if not label or _is_strain(label):
            label = f"{sheet}-{order + 1}"
        xu = _guess_unit(unit_row[xi] if unit_row and xi < len(unit_row) else "", "%")
        yu = _guess_unit(unit_row[yi] if unit_row and yi < len(unit_row) else "", "kPa")
        result.append(
            Series(
                key=f"{path}::{sheet}::{order}",
                label=label,
                group=sheet if sheet.upper() not in ("SHEET1", "") else os.path.basename(path),
                source=path,
                x=np.asarray(xs, float),
                y=np.asarray(ys, float),
                x_unit=xu,
                y_unit=yu,
                kind="xlsx",
            )
        )
    return result


# --------------------------------------------------------------------------- #
# 文件夹扫描
# --------------------------------------------------------------------------- #
def _natural_key(s: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def scan_folder(root: str, prefer: str = "auto", recursive: bool = True) -> ScanResult:
    """扫描文件夹，提取所有可用曲线。

    prefer: ``auto``（有 Excel 时只用 Excel）、``all``、``csv``、``xlsx``
    """
    res = ScanResult()
    if not root or not os.path.isdir(root):
        res.warnings.append(f"文件夹不存在：{root}")
        return res

    text_files: list[str] = []
    excel_files: list[str] = []
    iterator = os.walk(root) if recursive else [(root, [], os.listdir(root))]
    for dirpath, dirnames, filenames in iterator:
        dirnames[:] = sorted(
            [d for d in dirnames if not d.startswith(".") and d.lower() not in SKIP_DIR_HINTS],
            key=_natural_key,
        )
        for fn in sorted(filenames, key=_natural_key):
            if fn.startswith("~$") or fn.startswith("."):
                continue
            ext = os.path.splitext(fn)[1].lower()
            full = os.path.join(dirpath, fn)
            if ext in TEXT_EXT:
                text_files.append(full)
            elif ext in EXCEL_EXT:
                excel_files.append(full)

    def load_text(files: Iterable[str]) -> None:
        for f in files:
            name = os.path.basename(f).lower()
            if "results" in name or "summary" in name:
                # Bluehill 导出的"结果汇总表"只有几个统计点，不是曲线
                res.skipped.append(f"{os.path.basename(f)}: 统计汇总表，已跳过")
                continue
            try:
                for s in parse_text_table(f):
                    if s.n < 10:
                        res.skipped.append(f"{os.path.basename(f)}: 数据点过少（{s.n}），已跳过")
                        continue
                    res.series.append(s)
            except Exception as exc:  # noqa: BLE001 - 单个文件失败不影响整体
                res.skipped.append(f"{os.path.basename(f)}: {exc}")

    def load_excel(files: Iterable[str]) -> None:
        for f in files:
            try:
                res.series.extend(parse_excel(f))
            except Exception as exc:  # noqa: BLE001
                res.skipped.append(f"{os.path.basename(f)}: {exc}")

    mode = (prefer or "auto").lower()
    if mode == "csv":
        load_text(text_files)
    elif mode in ("xlsx", "excel"):
        load_excel(excel_files)
    elif mode == "all":
        load_excel(excel_files)
        load_text(text_files)
    else:  # auto
        if excel_files:
            load_excel(excel_files)
            if not res.series:
                load_text(text_files)
        else:
            load_text(text_files)

    if not res.series:
        res.warnings.append(
            "没有找到包含“拉伸应变 / 拉伸应力”两列的表格。\n"
            "请确认选择的是数据所在文件夹（如 *_RawData 文件夹或其上级）。"
        )
    res.series.sort(key=lambda s: (_natural_key(s.group), _natural_key(s.label)))
    return res


def folder_preview(root: str, limit: int = 400) -> list[str]:
    """列出文件夹内的数据文件（给界面显示用）。"""
    out: list[str] = []
    if not os.path.isdir(root):
        return out
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for fn in sorted(filenames, key=_natural_key):
            if os.path.splitext(fn)[1].lower() in TEXT_EXT | EXCEL_EXT:
                rel = os.path.relpath(os.path.join(dirpath, fn), root)
                out.append(rel)
                if len(out) >= limit:
                    return out
    return out
